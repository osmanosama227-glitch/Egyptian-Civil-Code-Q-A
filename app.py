# app.py — القانون المدني المصري: اسأل الموديل واحكم على إجابته
# Run on Colab (GPU) with the launcher notebook, or locally with a CUDA GPU:
#   streamlit run app.py --server.fileWatcherType none
#
# Needs next to this file:
#   egyptian_civil_code_articles2.json   (the articles)
#   egyptian_lora_model.zip              (from the training notebook)

# Unsloth must be imported before transformers / peft.
from unsloth import FastLanguageModel

import csv, datetime, glob, html, io, json, os, re, threading, zipfile
import numpy as np
import pandas as pd
import streamlit as st
import torch

ARTICLES_PATH  = "egyptian_civil_code_articles2.json"
ADAPTER_ZIP    = "egyptian_lora_model.zip"
ADAPTER_DIR    = "egyptian_lora_model"
LOG_PATH       = "validation_log.csv"
MAX_SEQ_LENGTH = 2048
TOP_N          = 50

st.set_page_config(page_title="القانون المدني المصري", page_icon="⚖️", layout="wide")
st.markdown("""
<style>
textarea {direction: rtl; text-align: right; font-size: 1.1rem;}
.rtl {direction: rtl; text-align: right; line-height: 2; font-size: 1.1rem;
      padding: .8rem 1rem; border-radius: .5rem; border: 1px solid rgba(128,128,128,.4);}
</style>
""", unsafe_allow_html=True)

# ------------------------------------------------------------------ Arabic text helpers
_DIACRITICS = re.compile(r"[\u0617-\u061A\u064B-\u0652\u0670\u0640]")
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")

def normalize_ar(text):
    t = _DIACRITICS.sub("", str(text))
    t = re.sub("[إأآٱ]", "ا", t)
    t = t.replace("ى", "ي").replace("ة", "ه").translate(_DIGITS)
    t = re.sub(r"[^\w\s]", " ", t)
    return " ".join(t.lower().split())

_PREFIXES = ("وال", "بال", "كال", "فال", "لل", "ال", "و", "ب", "ل", "ف", "ك")
def strip_prefix(w):
    for p in _PREFIXES:
        if w.startswith(p) and len(w) - len(p) >= 3:
            return w[len(p):]
    return w

def tokens(text):
    return [strip_prefix(w) for w in normalize_ar(text).split()]

def faithfulness(answer, context):
    """Share of the answer's words that also appear in the article (a rough check, not a proof)."""
    a = tokens(answer); c = set(tokens(context))
    return sum(t in c for t in a) / max(len(a), 1)

def rtl_box(text):
    return f'<div class="rtl">{html.escape(text).replace(chr(10), "<br>")}</div>'

def rrf_single(rank_lists, k=60):
    score = {}
    for r in rank_lists:
        for pos, idx in enumerate(r):
            score[int(idx)] = score.get(int(idx), 0.0) + 1.0 / (k + pos + 1)
    return sorted(score, key=score.get, reverse=True)

# ------------------------------------------------------------------ loading (once per session)
@st.cache_resource(show_spinner="جاري تحميل المواد وموديل البحث (مرة واحدة)...")
def load_retriever():
    from sentence_transformers import SentenceTransformer
    from rank_bm25 import BM25Okapi
    corpus = json.load(open(ARTICLES_PATH, encoding="utf-8"))
    embedder = SentenceTransformer("BAAI/bge-m3", device="cuda")
    embedder.max_seq_length = 1024
    embedder.half()
    texts = [a["context"] for a in corpus]
    emb = embedder.encode(texts, batch_size=16, normalize_embeddings=True,
                          convert_to_numpy=True, show_progress_bar=False).astype(np.float32)
    bm25 = BM25Okapi([tokens(t) for t in texts])
    art2idx = {a["article"]: i for i, a in enumerate(corpus)}
    return corpus, embedder, emb, bm25, art2idx

@st.cache_resource(show_spinner="جاري تحميل الموديل المتدرب (دقيقتين تقريباً)...")
def load_model():
    if not os.path.isdir(ADAPTER_DIR) and os.path.exists(ADAPTER_ZIP):
        zipfile.ZipFile(ADAPTER_ZIP).extractall(ADAPTER_DIR)
    cfgs = glob.glob(os.path.join(ADAPTER_DIR, "**", "adapter_config.json"), recursive=True)
    if not cfgs:
        raise FileNotFoundError(f"adapter_config.json not found. Put {ADAPTER_ZIP} next to app.py.")
    model, tokenizer = FastLanguageModel.from_pretrained(
        model_name=os.path.dirname(cfgs[0]), max_seq_length=MAX_SEQ_LENGTH,
        dtype=None, load_in_4bit=True)
    FastLanguageModel.for_inference(model)
    return model, tokenizer, threading.Lock()

try:
    corpus, embedder, c_emb, bm25, art2idx = load_retriever()
    model, tokenizer, gen_lock = load_model()
except Exception as e:
    st.error(f"تعذر تحميل الموديل أو الملفات: {e}")
    st.stop()

# ------------------------------------------------------------------ retrieval + generation
_ART_REF = re.compile(r"(?:الماده|ماده)\s*(?:رقم\s*)?(\d{1,4})")

def cited_articles(q):
    out = []
    for n in _ART_REF.findall(normalize_ar(q)):
        i = art2idx.get(int(n))
        if i is not None and i not in out:
            out.append(i)
    return out[:3]

def retrieve(question, k=3):
    q = embedder.encode([question], normalize_embeddings=True, convert_to_numpy=True,
                        show_progress_bar=False).astype(np.float32)[0]
    dense = c_emb @ q
    dense_rank = np.argsort(-dense)[:TOP_N]
    bm_rank = np.argsort(-bm25.get_scores(tokens(question)))[:TOP_N]
    fused = rrf_single([dense_rank, bm_rank])
    cited = cited_articles(question)
    order = cited + [i for i in fused if i not in cited]
    return order[:k], dense, cited

# Must be identical to the prompt used for training.
ALPACA = '''Below is a question about a law in Egypt (the Egyptian Civil Code, Law 131 of 1948), paired with an input that gives the text of the relevant article. Write an answer in formal Arabic that answers the question using only the article.

### Question:
{}

### Input:
{}

### Answer:
{}'''

def build_prompt(question, idxs):
    n = len(idxs)
    while True:
        ctx = "\n\n".join(corpus[i]["context"] for i in idxs[:n])
        prompt = ALPACA.format(question, ctx, "")
        if n == 1 or len(tokenizer(prompt)["input_ids"]) <= MAX_SEQ_LENGTH - 256:
            return prompt, list(idxs[:n]), ctx
        n -= 1

def generate(prompt, use_base=False):
    with gen_lock:
        inputs = tokenizer([prompt], return_tensors="pt").to("cuda")
        kw = dict(max_new_tokens=256, do_sample=False, use_cache=True,
                  pad_token_id=tokenizer.eos_token_id)
        with torch.no_grad():
            if use_base:
                with model.disable_adapter():
                    out = model.generate(**inputs, **kw)
            else:
                out = model.generate(**inputs, **kw)
        text = tokenizer.decode(out[0][inputs["input_ids"].shape[1]:], skip_special_tokens=True)
    return text.split("###")[0].strip()

def extract_text(uploaded):
    name, data = uploaded.name.lower(), uploaded.getvalue()
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    lines.extend(cell.text.split("\n"))
        return "\n".join(l for l in lines if l.strip())
    if name.endswith(".pdf"):
        import fitz
        with fitz.open(stream=data, filetype="pdf") as pdf:
            return "\n".join(p.get_text() for p in pdf)
    return ""

def save_rating(res, rating, note):
    new = not os.path.exists(LOG_PATH)
    with open(LOG_PATH, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["time", "question", "articles_given", "retrieved_top3", "mode",
                        "answer", "answer_in_article", "rating", "note"])
        w.writerow([res["time"], res["question"], res["used"], res["top3"], res["mode"],
                    res["answer"], round(res["faith"], 3), rating, note])

# ------------------------------------------------------------------ UI
MODE_TOP1, MODE_TOP3, MODE_NUM = ("أعلى نتيجة بحث (الأساسي)", "أعلى 3 نتايج (الموديل مدربش على كده)", "مادة بعينها برقمها")

with st.sidebar:
    st.header("الإعدادات")
    mode = st.radio("المادة اللي هتروح للموديل", [MODE_TOP1, MODE_TOP3, MODE_NUM])
    art_number = st.number_input("رقم المادة", min_value=1, max_value=1149, value=1) if mode == MODE_NUM else None
    compare = st.checkbox("قارن مع نفس الموديل من غير التدريب (base)")
    st.caption("أداة بحثية لتقييم النموذج، وليست بديلاً عن استشارة قانونية.")

st.title("⚖️ القانون المدني المصري — اسأل الموديل")
st.caption("اكتب سؤالك أو ارفع ملف Word/PDF فيه السؤال، وراجع الإجابة قصاد نص المادة.")

up = st.file_uploader("ملف فيه السؤال (اختياري)", type=["docx", "pdf"])
if up is not None:
    sig = (up.name, up.size)
    if st.session_state.get("_file_sig") != sig:
        txt = extract_text(up).strip()
        st.session_state["_file_sig"] = sig
        if txt:
            st.session_state["q_text"] = txt[:1500]
        else:
            st.warning("الملف مفيهوش نص يتقري (ممكن يكون صور ممسوحة). اكتب السؤال بإيدك.")

question = st.text_area("السؤال", key="q_text", height=130,
                        placeholder="مثال: متى يجوز للقاضي أن يعدل العقد؟")

if st.button("اسأل", type="primary"):
    if not question.strip():
        st.warning("اكتب سؤال الأول.")
    else:
        with st.spinner("بدور على المادة وبكتب الإجابة..."):
            order, dense, cited = retrieve(question)
            if mode == MODE_NUM:
                idxs = [art2idx[int(art_number)]] if int(art_number) in art2idx else []
            else:
                idxs = order[:1] if mode == MODE_TOP1 else order[:3]
            if not idxs:
                st.session_state["result"] = None
                st.error("المادة دي مش موجودة في الملف (ممكن تكون ملغاة).")
            else:
                prompt, used, ctx = build_prompt(question, idxs)
                answer = generate(prompt)
                st.session_state["result"] = {
                    "time": datetime.datetime.now().isoformat(timespec="seconds"),
                    "question": question, "mode": mode,
                    "used": [corpus[i]["article"] for i in used], "ctx": ctx,
                    "answer": answer, "faith": faithfulness(answer, ctx),
                    "base_answer": generate(prompt, use_base=True) if compare else None,
                    "top3": [corpus[i]["article"] for i in order[:3]],
                    "top3_info": [(corpus[i]["article"], float(dense[i]), i in cited, corpus[i]["context"]) for i in order[:3]],
                    "saved": False,
                }

res = st.session_state.get("result")
if res:
    st.divider()
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("إجابة الموديل")
        st.markdown(rtl_box(res["answer"] or "(الموديل مرجعش إجابة)"), unsafe_allow_html=True)
        pct = res["faith"] * 100
        st.metric("كلمات الإجابة الموجودة في المادة", f"{pct:.0f}%")
        if pct < 60:
            st.warning("نسبة منخفضة: ممكن الإجابة فيها حاجة مش في المادة. راجعها بدقة.")
        if res["base_answer"] is not None:
            st.subheader("نفس الموديل من غير تدريب (base)")
            st.markdown(rtl_box(res["base_answer"] or "(فاضي)"), unsafe_allow_html=True)
    with c2:
        st.subheader("المادة اللي اتبعتت للموديل: " + "، ".join(str(a) for a in res["used"]))
        st.markdown(rtl_box(res["ctx"]), unsafe_allow_html=True)

    with st.expander("أعلى 3 نتايج بحث (لو المادة الصح مش هنا يبقى المشكلة في البحث مش في الموديل)"):
        for art, score, was_cited, text in res["top3_info"]:
            tag = " — اتذكرت برقمها في السؤال" if was_cited else f" — تشابه {score:.2f}"
            st.markdown(f"**المادة {art}**{tag}")
            st.markdown(rtl_box(text), unsafe_allow_html=True)

    st.divider()
    st.subheader("تقييمك")
    rating = st.radio("الإجابة دي:", ["✅ صحيحة ومدعومة بالمادة", "⚠️ صحيحة جزئياً أو ناقصة",
                                      "❌ خاطئة أو فيها حاجة مش في المادة",
                                      "🔎 المادة المسترجعة غلط (مشكلة بحث)"], key="rating")
    note = st.text_input("ملاحظة (اختياري)", key="note")
    if st.button("حفظ التقييم", disabled=res["saved"]):
        save_rating(res, rating, note)
        res["saved"] = True
        st.rerun()
    if res["saved"]:
        st.success("اتحفظ ✅")

if os.path.exists(LOG_PATH):
    log = pd.read_csv(LOG_PATH, encoding="utf-8-sig")
    with st.sidebar:
        st.divider()
        st.write(f"التقييمات المحفوظة: **{len(log)}**")
        st.write(log["rating"].value_counts().to_string())
        st.download_button("نزّل سجل التقييمات (CSV)", log.to_csv(index=False, encoding="utf-8-sig"),
                           file_name="validation_log.csv", mime="text/csv")
